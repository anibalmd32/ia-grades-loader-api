"""Descarga de notas en Excel (.xlsx) para el profesor.

- Materia completa: hoja «Planilla» con la nota de cada lapso, el promedio, la definitiva y
  la situación, y una hoja más por cada lapso con plan de evaluación.
- Un lapso: sus evaluaciones (con porcentaje), lo acumulado y la nota del lapso.

Todos los números salen de los mismos cálculos que la API (GradingRules y
services.evaluations), así que el Excel coincide con lo que se ve en la app.
"""

import io
import re
import unicodedata
from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.models import Evaluation, Subject
from app.services.evaluations import EvaluationRow
from app.services.grades import StudentRow
from app.services.rules import GradingRules, fmt

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FILL = PatternFill("solid", fgColor="DCEBFA")
_TABLE_START = 4  # fila de los encabezados (las anteriores: título y reglas)
# Excel no admite estos caracteres en el nombre de una hoja y lo limita a 31
_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")


def _number_format(rules: GradingRules) -> str:
    return "0" if rules.grade_decimals == 0 else "0." + "0" * rules.grade_decimals


def _value(value: Decimal | None) -> float | None:
    return None if value is None else float(value)


def _standing(row: StudentRow) -> str:
    if row.passed is None:
        return "En curso"
    return "Aprobado" if row.passed else "Reprobado"


def _sheet_title(text: str) -> str:
    return _SHEET_FORBIDDEN.sub("-", text)[:31]


def _write_table(
    sheet: Worksheet,
    title: str,
    subtitle: str,
    headers: list[str],
    rows: list[list],
    numeric_from: int,
    number_format: str,
) -> None:
    """Título, reglas, encabezados en negrita, filtros, paneles fijos y anchos de columna.
    Las columnas desde `numeric_from` (0-based) son notas con el formato de la escala."""
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(italic=True, color="555555")
    sheet.append([])
    sheet.append(headers)
    _as_text(sheet[1] + sheet[2] + sheet[_TABLE_START])
    for cell in sheet[_TABLE_START]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for values in rows:
        sheet.append(values)
        _as_text(sheet[sheet.max_row])
        for cell in sheet[sheet.max_row][numeric_from:]:
            if isinstance(cell.value, float):
                cell.number_format = number_format
                cell.alignment = Alignment(horizontal="center")
    last_column = get_column_letter(len(headers))
    sheet.auto_filter.ref = f"A{_TABLE_START}:{last_column}{max(sheet.max_row, _TABLE_START)}"
    sheet.freeze_panes = f"D{_TABLE_START + 1}"
    for index, header in enumerate(headers, start=1):
        widest = max(
            [len(str(header))] + [len(str(r[index - 1])) for r in rows if r[index - 1] is not None]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(max(widest + 2, 10), 40)


def _as_text(cells) -> None:
    """Los textos (nombres, materias, evaluaciones) nunca son fórmulas.

    openpyxl guarda como fórmula cualquier texto que empiece por "=": un apellido como
    `=HYPERLINK("https://...?"&A5, "Ver")` se ejecutaría al abrir el Excel del profesor.
    """
    for cell in cells:
        if cell.data_type == "f":
            cell.data_type = "s"


def _person(student) -> list:
    return [student.cedula_display or "", student.last_name, student.first_name]


def _person_headers(document_label: str) -> list[str]:
    return [document_label, "Apellido", "Nombre"]


def _heading(subject: Subject, rules: GradingRules) -> str:
    return f"{subject.code} · {subject.name} — {rules.period_label} {subject.semester.name}"


def _generated() -> str:
    return datetime.now(UTC).strftime("%d/%m/%Y %H:%M UTC")


def _term_sheet(
    sheet: Worksheet,
    subject: Subject,
    rules: GradingRules,
    term: int,
    plan: list[Evaluation],
    rows: list[EvaluationRow],
    document_label: str,
) -> None:
    term_name = rules.term_name(term)
    total = sum((e.weight for e in plan), Decimal(0))
    headers = [
        *_person_headers(document_label),
        *(f"{e.name} ({fmt(e.weight)} %)" for e in plan),
        "Acumulado",
        term_name,
    ]
    _write_table(
        sheet,
        f"{_heading(subject, rules)} · {term_name}",
        f"Plan de evaluación al {fmt(total)} %. La nota del {term_name.lower()} es la suma de "
        f"cada nota por su porcentaje y solo aparece con el plan al 100 % y todo calificado. "
        f"Notas {rules.scale_text}. Generado el {_generated()}.",
        headers,
        [
            [
                *_person(row.student),
                *(_value(row.values.get(e.id)) for e in plan),
                _value(row.accumulated),
                _value(row.term_grade),
            ]
            for row in rows
        ],
        numeric_from=3,
        number_format=_number_format(rules),
    )


def _save(workbook: Workbook) -> bytes:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def subject_workbook(
    subject: Subject,
    rules: GradingRules,
    roster: list[StudentRow],
    term_plans: dict[int, tuple[list[Evaluation], list[EvaluationRow]]],
    document_label: str,
) -> bytes:
    """Planilla completa de la materia y una hoja por cada lapso con plan de evaluación."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Planilla"
    headers = [
        *_person_headers(document_label),
        *(rules.term_name(t) for t in rules.terms),
        "Promedio",
        "Definitiva",
        "Situación",
    ]
    _write_table(
        sheet,
        _heading(subject, rules),
        f"{rules.final_grade_text}. Notas {rules.scale_text}. Generado el {_generated()}.",
        headers,
        [
            [
                *_person(row.student),
                *(_value(row.grades.get(t)) for t in rules.terms),
                _value(row.average),
                _value(row.final_grade),
                _standing(row),
            ]
            for row in roster
        ],
        numeric_from=3,
        number_format=_number_format(rules),
    )
    for term, (plan, rows) in sorted(term_plans.items()):
        _term_sheet(
            workbook.create_sheet(_sheet_title(rules.term_name(term))),
            subject,
            rules,
            term,
            plan,
            rows,
            document_label,
        )
    return _save(workbook)


def term_workbook(
    subject: Subject,
    rules: GradingRules,
    term: int,
    plan: list[Evaluation],
    rows: list[EvaluationRow],
    document_label: str,
) -> bytes:
    """Evaluaciones de un lapso con lo acumulado y la nota del lapso."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _sheet_title(rules.term_name(term))
    _term_sheet(sheet, subject, rules, term, plan, rows, document_label)
    return _save(workbook)


def _slug(text: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    return re.sub(r"[^A-Za-z0-9]+", "-", ascii_text).strip("-").lower() or "notas"


def file_name(subject: Subject, rules: GradingRules, term: int | None = None) -> str:
    """notas-mat-102-2026-i.xlsx o notas-mat-102-2026-i-lapso-1.xlsx"""
    parts = [subject.code, subject.semester.name]
    if term is not None:
        parts.append(rules.term_name(term))
    return f"notas-{_slug(' '.join(parts))}.xlsx"


def content_disposition(name: str) -> str:
    return f'attachment; filename="{name}"'
