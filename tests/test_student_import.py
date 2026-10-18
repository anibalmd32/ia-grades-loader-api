import io

import pytest
from openpyxl import Workbook, load_workbook

from app.models import Role


def xlsx(rows: list[list]) -> bytes:
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture
async def coord(login_as):
    return await login_as(Role.COORDINATOR, email="coordina@uni.edu")


async def upload(client, headers, content: bytes, dry_run: bool):
    return await client.post(
        "/students/import",
        files={"file": ("estudiantes.xlsx", content, "application/octet-stream")},
        data={"dry_run": str(dry_run).lower()},
        headers=headers,
    )


async def test_plantilla(client, coord):
    response = await client.get("/students/import/template", headers=coord)
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert [c.value for c in sheet[1]] == ["Nombre", "Apellido", "Tipo", "Cédula", "Materias"]


async def test_validar_y_confirmar(client, coord, campus):
    content = xlsx(
        [
            ["Listado de estudiantes 2026"],
            ["Nombres", "Apellidos", "Tipo", "C.I.", "Materias"],
            ["Luis", "Gómez", "V", "V-20.111.222", "MAT-101, fis-101"],
            ["Mary", "Smith", "E", "81234567", ""],
            ["Ana", "Repetida", "V", "1", "MAT-201"],  # ya existe (Ana del campus)
            ["Sin", "Cedula", "V", "", ""],
            ["Otra", "Vez", "V", "20111222", ""],  # repetida en el archivo
            ["Pepe", "", "V", "333", "QUI-999"],  # sin apellido y materia inexistente
        ]
    )
    preview = await upload(client, coord, content, dry_run=True)
    assert preview.status_code == 200, preview.text
    body = preview.json()
    assert body["dry_run"] is True
    assert body["summary"] == {"new": 2, "existing": 1, "error": 3}
    luis = body["rows"][0]
    assert (luis["row"], luis["nationality"], luis["cedula"]) == (3, "V", 20111222)
    assert luis["subject_codes"] == ["FIS-101", "MAT-101"]
    assert luis["generated_password"] is None
    errors = {r["row"]: r["issues"] for r in body["rows"] if r["status"] == "error"}
    assert errors[6] == ["Falta la cédula"]
    assert "repetida en el archivo (fila 3)" in errors[7][0]
    assert set(errors[8]) == {"Falta el apellido", "La materia QUI-999 no existe"}
    # Nada creado todavía
    total = (await client.get("/students", headers=coord)).json()["total"]
    assert total == 2

    done = (await upload(client, coord, content, dry_run=False)).json()
    assert done["summary"] == {"new": 2, "existing": 1, "error": 3}
    created = [r for r in done["rows"] if r["status"] == "new"]
    assert all(len(r["generated_password"]) >= 8 for r in created)
    assert created[0]["user"]["must_change_password"] is True
    assert (await client.get("/students", headers=coord)).json()["total"] == 4

    luis_id = created[0]["user"]["id"]
    subjects = (await client.get(f"/students/{luis_id}/subjects", headers=coord)).json()
    assert sorted(s["code"] for s in subjects) == ["FIS-101", "MAT-101"]
    ana = campus["students"]["Ana"]["id"]
    ana_subjects = (await client.get(f"/students/{ana}/subjects", headers=coord)).json()
    assert len(ana_subjects) == 3  # MAT-201 ya la tenía: no se duplica

    # El estudiante nuevo puede entrar con su contraseña
    login = await client.post(
        "/auth/login/student",
        json={"nationality": "V", "cedula": 20111222, "password": created[0]["generated_password"]},
    )
    assert login.status_code == 200


async def test_columna_de_nombre_completo(client, coord):
    content = xlsx([["Estudiante", "Cédula"], ["Pérez Gómez, Ana María", "E-5"]])
    row = (await upload(client, coord, content, dry_run=True)).json()["rows"][0]
    assert (row["first_name"], row["last_name"], row["nationality"]) == (
        "Ana María",
        "Pérez Gómez",
        "E",
    )


async def test_archivo_sin_encabezados(client, coord):
    response = await upload(client, coord, xlsx([["a", "b"], ["c", "d"]]), dry_run=True)
    assert response.status_code == 400
    assert "encabezados" in response.json()["detail"]


async def test_periodo_cerrado(client, coord, campus):
    mat = campus["subjects"]["MAT-101"]
    await client.post(f"/semesters/{mat['semester']['id']}/close", headers=coord)
    content = xlsx(
        [["Nombre", "Apellido", "Cédula", "Materias"], ["Luis", "Gómez", "9", "MAT-101"]]
    )
    row = (await upload(client, coord, content, dry_run=True)).json()["rows"][0]
    assert row["issues"] == ["La materia MAT-101 es de un periodo cerrado"]


async def test_solo_coordinacion(client, campus):
    teacher = campus["teachers"]["Pedro Profe"]
    assert (await client.get("/students/import/template", headers=teacher)).status_code == 403
