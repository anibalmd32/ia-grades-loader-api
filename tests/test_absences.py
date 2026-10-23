"""Inasistencias: el estudiante que no se presenta cuenta con la nota mínima de la escala."""

import pytest

from app.services.grade_import import is_absent_mark
from tests.conftest import active_token


async def put_grades(client, course, entries):
    return await client.put(
        f"{course['url']}/grades", json={"grades": entries}, headers=course["headers"]
    )


async def put_evaluation_grades(client, course, term, entries):
    return await client.put(
        f"{course['url']}/terms/{term}/evaluation-grades",
        json={"grades": entries},
        headers=course["headers"],
    )


async def test_inasistente_en_un_lapso(client, course):
    ana = course["students"]["Ana Pérez"]
    response = await put_grades(
        client,
        course,
        [
            {"student_id": ana, "term": 1, "absent": True},
            {"student_id": ana, "term": 2, "value": 16},
            {"student_id": ana, "term": 3, "value": 13},
        ],
    )
    assert response.status_code == 200, response.text
    row = next(s for s in response.json()["students"] if s["student"]["id"] == ana)
    assert row["grades"][0] == {"term": 1, "value": 1, "absent": True}
    assert row["grades"][1]["absent"] is False
    # Cuenta con la nota mínima: (1 + 16 + 13) / 3
    assert row["final_grade"] == 10
    assert row["passed"] is True

    # Corregir la inasistencia con una nota la quita
    response = await put_grades(client, course, [{"student_id": ana, "term": 1, "value": 14}])
    row = next(s for s in response.json()["students"] if s["student"]["id"] == ana)
    assert row["grades"][0] == {"term": 1, "value": 14, "absent": False}


async def test_inasistente_no_lleva_nota(client, course):
    ana = course["students"]["Ana Pérez"]
    response = await put_grades(
        client, course, [{"student_id": ana, "term": 1, "value": 12, "absent": True}]
    )
    assert response.status_code == 422


async def test_historial_y_aviso_de_la_inasistencia(client, course, session):
    ana = course["students"]["Ana Pérez"]
    await put_grades(client, course, [{"student_id": ana, "term": 1, "value": 9}])
    await put_grades(client, course, [{"student_id": ana, "term": 1, "absent": True}])
    # Repetirla no cambia nada
    await put_grades(client, course, [{"student_id": ana, "term": 1, "absent": True}])

    history = (await client.get(f"{course['url']}/grade-history", headers=course["headers"])).json()
    assert [
        (c["old_value"], c["old_absent"], c["new_value"], c["new_absent"]) for c in history["items"]
    ] == [(9, False, 1, True), (None, False, 9, False)]

    token = await active_token(session, ana)
    headers = {"Authorization": f"Bearer {token}"}
    notices = (await client.get("/student/notifications", headers=headers)).json()["items"]
    assert notices[0]["body"] == "Lapso 1: inasistente"
    subject = (await client.get("/student/subjects", headers=headers)).json()["subjects"][0]
    assert subject["grades"][0] == {"term": 1, "value": 1, "absent": True}


@pytest.fixture
async def plan(client, course):
    """Lapso 1 de MAT-101 con dos evaluaciones al 50 %."""
    ids = []
    for name in ("Parcial", "Taller"):
        response = await client.post(
            f"{course['url']}/evaluations",
            json={"term": 1, "name": name, "weight": 50},
            headers=course["headers"],
        )
        ids.append(response.json()["id"])
    return ids


async def test_inasistente_en_una_evaluacion(client, course, plan, session):
    ana = course["students"]["Ana Pérez"]
    parcial, taller = plan
    response = await put_evaluation_grades(
        client,
        course,
        1,
        [
            {"student_id": ana, "evaluation_id": parcial, "absent": True},
            {"student_id": ana, "evaluation_id": taller, "value": 19},
        ],
    )
    assert response.status_code == 200, response.text
    row = next(s for s in response.json()["students"] if s["student"]["id"] == ana)
    assert row["grades"] == [
        {"evaluation_id": parcial, "value": 1, "absent": True},
        {"evaluation_id": taller, "value": 19, "absent": False},
    ]
    assert (row["term_grade"], row["term_absent"]) == (10, False)

    token = await active_token(session, ana)
    headers = {"Authorization": f"Bearer {token}"}
    subject = (await client.get("/student/subjects", headers=headers)).json()["subjects"][0]
    assert [(e["name"], e["absent"]) for e in subject["evaluations"]] == [
        ("Parcial", True),
        ("Taller", False),
    ]
    notice = (await client.get("/student/notifications", headers=headers)).json()["items"][0]
    assert "Parcial (Lapso 1): inasistente" in notice["body"]


async def test_inasistente_en_todo_un_lapso_con_plan(client, course, plan):
    ana = course["students"]["Ana Pérez"]
    # En un lapso con plan no se carga la nota a mano...
    blocked = await put_grades(client, course, [{"student_id": ana, "term": 1, "value": 15}])
    assert blocked.status_code == 422
    # ...pero sí se marca la inasistencia en todo el lapso: queda en cada evaluación
    response = await put_grades(client, course, [{"student_id": ana, "term": 1, "absent": True}])
    assert response.status_code == 200, response.text
    row = next(s for s in response.json()["students"] if s["student"]["id"] == ana)
    assert row["grades"][0] == {"term": 1, "value": 1, "absent": True}

    sheet = (
        await client.get(f"{course['url']}/terms/1/evaluation-grades", headers=course["headers"])
    ).json()
    row = next(s for s in sheet["students"] if s["student"]["id"] == ana)
    assert all(g["absent"] for g in row["grades"])
    assert row["term_absent"] is True

    # Si luego se presenta a una evaluación, el lapso deja de ser inasistente
    response = await put_evaluation_grades(
        client, course, 1, [{"student_id": ana, "evaluation_id": plan[1], "value": 20}]
    )
    row = next(s for s in response.json()["students"] if s["student"]["id"] == ana)
    assert (row["term_grade"], row["term_absent"]) == (10.5, False)


async def test_exportar_muestra_la_inasistencia(client, course):
    from io import BytesIO

    from openpyxl import load_workbook

    ana = course["students"]["Ana Pérez"]
    await put_grades(client, course, [{"student_id": ana, "term": 2, "absent": True}])
    response = await client.get(f"{course['url']}/grades/export", headers=course["headers"])
    sheet = load_workbook(BytesIO(response.content)).active
    rows = list(sheet.iter_rows(min_row=5, values_only=True))
    ana_row = next(r for r in rows if r[1] == "Pérez")
    assert ana_row[4] == "Inasistente"


@pytest.mark.parametrize(
    ("raw", "expected"),
    [("I", True), ("NP", True), (" inasistente ", True), ("No se presentó", True), ("15", False)],
)
def test_la_importacion_reconoce_las_inasistencias(raw, expected):
    assert is_absent_mark(raw) is expected
    assert is_absent_mark(15) is False
