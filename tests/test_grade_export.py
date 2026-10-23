import io

from openpyxl import load_workbook

from app.models import Role

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def table(sheet):
    """Encabezados (fila 4) y filas de datos como listas de valores."""
    rows = [list(r) for r in sheet.iter_rows(min_row=4, values_only=True)]
    return rows[0], rows[1:]


async def download(client, url, headers):
    response = await client.get(url, headers=headers)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX
    return response, load_workbook(io.BytesIO(response.content))


async def test_descarga_la_planilla_de_la_materia(client, campus):
    pedro = campus["teachers"]["Pedro Profe"]
    response, workbook = await download(client, f"{campus['urls']['MAT-101']}/grades/export", pedro)
    assert (
        'filename="notas-mat-101-2026-i-seccion-a.xlsx"' in response.headers["content-disposition"]
    )
    assert workbook.sheetnames == ["Planilla"]
    sheet = workbook["Planilla"]
    assert sheet["A1"].value == "MAT-101 · Mate I — Sección A — Semestre 2026-I"
    headers, rows = table(sheet)
    assert headers == [
        "Cédula", "Apellido", "Nombre", "Lapso 1", "Lapso 2", "Lapso 3",
        "Promedio", "Definitiva", "Situación",
    ]  # fmt: skip
    by_name = {r[2]: r for r in rows}
    # Ana: 15 y 18.5 (promedio 16.75, sin definitiva); Beto: 20
    assert by_name["Ana"][:9] == ["V-1", "Test", "Ana", 15, 18.5, None, 16.75, None, "En curso"]
    assert by_name["Beto"][3] == 20
    assert sheet.cell(row=5, column=4).number_format == "0.00"


async def test_incluye_una_hoja_por_lapso_con_plan(client, campus):
    pedro = campus["teachers"]["Pedro Profe"]
    ana = campus["students"]["Ana"]["id"]
    examen = (
        await client.post(
            f"{campus['urls']['MAT-101']}/evaluations",
            json={"term": 3, "name": "Examen", "weight": 60},
            headers=pedro,
        )
    ).json()
    await client.post(
        f"{campus['urls']['MAT-101']}/evaluations",
        json={"term": 3, "name": "Taller", "weight": 40},
        headers=pedro,
    )
    await client.put(
        f"{campus['urls']['MAT-101']}/terms/3/evaluation-grades",
        json={"grades": [{"student_id": ana, "evaluation_id": examen["id"], "value": 15}]},
        headers=pedro,
    )
    _, workbook = await download(client, f"{campus['urls']['MAT-101']}/grades/export", pedro)
    assert workbook.sheetnames == ["Planilla", "Lapso 3"]
    headers, rows = table(workbook["Lapso 3"])
    assert headers == [
        "Cédula", "Apellido", "Nombre", "Examen (60 %)", "Taller (40 %)", "Acumulado", "Lapso 3",
    ]  # fmt: skip
    ana_row = next(r for r in rows if r[2] == "Ana")
    assert ana_row[3:] == [15, None, 9, None]

    # El mismo lapso por separado
    response, single = await download(
        client, f"{campus['urls']['MAT-101']}/terms/3/evaluation-grades/export", pedro
    )
    assert (
        'filename="notas-mat-101-2026-i-seccion-a-lapso-3.xlsx"'
        in response.headers["content-disposition"]
    )
    assert single.sheetnames == ["Lapso 3"]
    assert "Plan de evaluación al 100 %" in single["Lapso 3"]["A2"].value
    assert table(single["Lapso 3"]) == (headers, rows)


async def test_lapso_sin_plan_y_validaciones(client, campus, login_as):
    pedro = campus["teachers"]["Pedro Profe"]
    _, workbook = await download(
        client, f"{campus['urls']['MAT-101']}/terms/1/evaluation-grades/export", pedro
    )
    headers, rows = table(workbook.active)
    assert headers == ["Cédula", "Apellido", "Nombre", "Acumulado", "Lapso 1"]
    assert len(rows) == 2

    bad_term = await client.get(
        f"{campus['urls']['MAT-101']}/terms/9/evaluation-grades/export", headers=pedro
    )
    assert bad_term.status_code == 422
    other = await login_as(Role.TEACHER, email="otro@uni.edu")
    assert (
        await client.get(f"{campus['urls']['MAT-101']}/grades/export", headers=other)
    ).status_code == 404
    student = campus["students"]["Ana"]["headers"]
    assert (
        await client.get(f"{campus['urls']['MAT-101']}/grades/export", headers=student)
    ).status_code == 403


async def test_el_front_puede_leer_el_nombre_del_archivo(client, campus):
    response = await client.get(
        f"{campus['urls']['MAT-101']}/grades/export",
        headers={**campus["teachers"]["Pedro Profe"], "Origin": "http://localhost:1420"},
    )
    exposed = response.headers["access-control-expose-headers"].lower()
    assert "content-disposition" in exposed
